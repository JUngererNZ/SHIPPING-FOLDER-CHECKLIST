import json
import os
import openpyxl
from openpyxl.styles import PatternFill

def populate_ops_checklist(source_template_xlsx: str, output_xlsx: str):
    if not os.path.exists(source_template_xlsx):
        print(f"[Error] Excel template source target file not found: '{source_template_xlsx}'")
        return

    # Load data gathered during the previous RAG step
    with open("structured_metadata.json", "r", encoding="utf-8") as f:
        extracted_data = json.load(f)

    # Pool verified values from files into a singular operational overlay dictionary
    master_field_map = {}
    for path, fields in extracted_data.items():
        for key, value in fields.items():
            if value and "unknown" not in str(value).lower() and str(value).strip() != "":
                master_field_map[key] = str(value).strip()

    print("Consolidated Operational Parameters Found:")
    print(json.dumps(master_field_map, indent=2))

    print(f"\nOpening checklist layout model workbook: {source_template_xlsx}")
    wb = openpyxl.load_workbook(source_template_xlsx)
    ws = wb.active # Accesses active grid worksheet

    # Coordinate mapping dict connecting labels to the exact row locations in Column B
    field_row_locations = {
        "Client Ref": 2,    # B2
        "Consignee": 3,     # B3
        "Description": 4,   # B4
        "PIN No": 5,        # B5
        "Serial No": 6,     # B6
        "Vessel": 7,        # B7
        "Voy": 8,           # B8
        "Tariff Code": 11   # B11 ("Tarrif Code, Weight & Cube")
    }

    # Hex codes matching standard traffic/robot colors
    amber_fill = PatternFill(start_color="FFBF00", end_color="FFBF00", fill_type="solid")  # Action Required Alert
    green_fill = PatternFill(start_color="C1E1C1", end_color="C1E1C1", fill_type="solid")  # Verified/Populated

    print("\nExecuting spreadsheet grid alignment updates...")
    for field_name, row_num in field_row_locations.items():
        cell = ws.cell(row=row_num, column=2) # Targeting Column B
        current_cell_val = str(cell.value or '').strip()

        # Route 1: Target field was successfully discovered by the RAG step
        if field_name in master_field_map:
            cell.value = master_field_map[field_name]
            cell.fill = green_fill
            print(f" [COMPLETE] Row {row_num} - {field_name} assigned: '{cell.value}' (Green)")
            
        # Route 2: Target field missing from RAG, check if cell contains pre-filled data
        else:
            if current_cell_val == "" or current_cell_val == "None":
                cell.value = ""
                cell.fill = amber_fill
                print(f" [ACTION REQUIRED] Row {row_num} - {field_name} is blank. Flagged AMBER.")
            else:
                cell.fill = green_fill
                print(f" [EXISTING DATA] Row {row_num} - {field_name} has pre-existing value: '{current_cell_val}' (Green)")

    # Save out as production workbook file
    wb.save(output_xlsx)
    print(f"\nVerification complete! Operational master file updated: '{output_xlsx}'")

if __name__ == "__main__":
    # Ensure FILE CHECKLIST.xlsx is located in your execution working folder directory
    INPUT_TEMPLATE = "FILE CHECKLIST.xlsx" 
    OUTPUT_PRODUCTION_SHEET = "PRODUCTION_FILE_CHECKLIST.xlsx"
    
    populate_ops_checklist(INPUT_TEMPLATE, OUTPUT_PRODUCTION_SHEET)